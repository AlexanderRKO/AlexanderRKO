"""Export a parsed MYE file to CSV, Excel, JSON and QuickBooks IIF, and
re-import edited CSVs so they can be packed back into a .MYE."""

from __future__ import annotations

import csv
import json
import os
from decimal import Decimal, InvalidOperation
from typing import List

from .core import Account, JournalEntry, JournalLine, MyeFile

ACCOUNTS_CSV = "accounts.csv"
JOURNAL_CSV = "journal.csv"
TRIAL_BALANCE_CSV = "trial_balance.csv"
COMPANY_CSV = "company.csv"
EXTRACT_INF = "Extract.inf"

JOURNAL_HEADER = [
    "entry",
    "date",
    "reference",
    "account_code",
    "account_name",
    "debit",
    "credit",
    "amount",
    "memo",
]


def _writer(path: str):
    return open(path, "w", newline="", encoding="utf-8-sig")


# ----- CSV ---------------------------------------------------------------


# Friendly names for the six fields of the [MYOB2000.05] company line,
# by position. Used so the editable company.csv preserves every field
# (including the ABN, which lives in position 2) with no data loss.
COMPANY_FIELD_NAMES = [
    "company_name",
    "address",
    "abn",
    "company_field_4",
    "period_start",
    "period_end",
]


def export_company_csv(mye: MyeFile, path: str) -> None:
    info = mye.extract_info()
    with _writer(path) as fh:
        w = csv.writer(fh)
        w.writerow(["field", "value"])
        for i, value in enumerate(mye.company_fields):
            name = COMPANY_FIELD_NAMES[i] if i < len(COMPANY_FIELD_NAMES) else f"company_field_{i}"
            w.writerow([name, value])
        # Serialisation quirks, so unpack -> pack is byte-exact too.
        term = "CRCRLF" if mye.journal_line_terminator == "\r\r\n" else "CRLF"
        w.writerow(["meta.journal_line_terminator", term])
        w.writerow(["meta.trailing_blank", "yes" if mye.trailing_blank else "no"])
        for key, value in info.items():
            w.writerow([f"extract.{key}", value])


def export_accounts_csv(mye: MyeFile, path: str) -> None:
    with _writer(path) as fh:
        w = csv.writer(fh)
        w.writerow(["account_code", "account_name"])
        for account in mye.accounts:
            w.writerow([account.code, account.name])


def export_journal_csv(mye: MyeFile, path: str) -> None:
    names = mye.account_map()
    with _writer(path) as fh:
        w = csv.writer(fh)
        w.writerow(JOURNAL_HEADER)
        for i, entry in enumerate(mye.entries, 1):
            for line in entry.lines:
                w.writerow(
                    [
                        i,
                        line.date,
                        line.reference,
                        line.account_code,
                        names.get(line.account_code, ""),
                        f"{line.debit:.2f}" if line.debit else "",
                        f"{line.credit:.2f}" if line.credit else "",
                        line.amount_text,
                        line.memo,
                    ]
                )


def export_trial_balance_csv(mye: MyeFile, path: str) -> None:
    with _writer(path) as fh:
        w = csv.writer(fh)
        w.writerow(["account_code", "account_name", "debits", "credits", "net_movement"])
        for code, name, debit, credit, net in mye.trial_balance():
            w.writerow([code, name, f"{debit:.2f}", f"{credit:.2f}", f"{net:.2f}"])


def export_csv(mye: MyeFile, out_dir: str) -> List[str]:
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    for name, fn in (
        (COMPANY_CSV, export_company_csv),
        (ACCOUNTS_CSV, export_accounts_csv),
        (JOURNAL_CSV, export_journal_csv),
        (TRIAL_BALANCE_CSV, export_trial_balance_csv),
    ):
        path = os.path.join(out_dir, name)
        fn(mye, path)
        paths.append(path)
    return paths


# ----- Excel -------------------------------------------------------------


def export_xlsx(mye: MyeFile, path: str) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "Excel export needs the 'openpyxl' package: pip install openpyxl"
        ) from exc

    bold = Font(bold=True)
    money = "#,##0.00"
    wb = Workbook()

    def style_header(ws):
        for cell in ws[1]:
            cell.font = bold
        ws.freeze_panes = "A2"

    def autosize(ws, widths):
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

    ws = wb.active
    ws.title = "Company"
    ws.append(["Field", "Value"])
    ws.append(["Company name", mye.company_name])
    ws.append(["Address", mye.address])
    ws.append(["Period start", mye.period_start])
    ws.append(["Period end", mye.period_end])
    for key, value in mye.extract_info().items():
        ws.append([key, value])
    style_header(ws)
    autosize(ws, [24, 50])

    ws = wb.create_sheet("Accounts")
    ws.append(["Code", "Account name"])
    for account in mye.accounts:
        ws.append([account.code, account.name])
    style_header(ws)
    autosize(ws, [10, 50])

    ws = wb.create_sheet("Journal")
    ws.append(["Entry", "Date", "Ref", "Account", "Account name", "Debit", "Credit", "Memo"])
    names = mye.account_map()
    for i, entry in enumerate(mye.entries, 1):
        for line in entry.lines:
            row = [
                i,
                line.date,
                line.reference,
                line.account_code,
                names.get(line.account_code, ""),
                float(line.debit) if line.debit else None,
                float(line.credit) if line.credit else None,
                line.memo,
            ]
            ws.append(row)
            ws.cell(row=ws.max_row, column=6).number_format = money
            ws.cell(row=ws.max_row, column=7).number_format = money
    style_header(ws)
    autosize(ws, [8, 12, 8, 10, 32, 14, 14, 50])

    ws = wb.create_sheet("Trial Balance")
    ws.append(["Code", "Account name", "Debits", "Credits", "Net movement"])
    for code, name, debit, credit, net in mye.trial_balance():
        ws.append([code, name, float(debit), float(credit), float(net)])
        for col in (3, 4, 5):
            ws.cell(row=ws.max_row, column=col).number_format = money
    style_header(ws)
    autosize(ws, [10, 40, 14, 14, 14])

    wb.save(path)
    return path


# ----- JSON ----------------------------------------------------------------


def export_json(mye: MyeFile, path: str) -> str:
    names = mye.account_map()
    doc = {
        "company": {
            "name": mye.company_name,
            "address": mye.address,
            "period_start": mye.period_start,
            "period_end": mye.period_end,
        },
        "extract_info": mye.extract_info(),
        "accounts": [{"code": a.code, "name": a.name} for a in mye.accounts],
        "journal": [
            {
                "reference": entry.reference,
                "date": entry.date,
                "lines": [
                    {
                        "date": line.date,
                        "account_code": line.account_code,
                        "account_name": names.get(line.account_code, ""),
                        "amount": str(line.amount),
                        "memo": line.memo,
                    }
                    for line in entry.lines
                ],
            }
            for entry in mye.entries
        ],
    }
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(doc, fh, indent=2, ensure_ascii=False)
    return path


# ----- QuickBooks IIF -------------------------------------------------------


def export_iif(mye: MyeFile, path: str) -> str:
    """General-journal IIF accepted by QuickBooks Desktop."""
    names = mye.account_map()

    def acct(code: str) -> str:
        name = names.get(code, "")
        return f"{code} {name}".strip().replace("\t", " ")

    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write("!TRNS\tTRNSTYPE\tDATE\tACCNT\tAMOUNT\tDOCNUM\tMEMO\n")
        fh.write("!SPL\tTRNSTYPE\tDATE\tACCNT\tAMOUNT\tDOCNUM\tMEMO\n")
        fh.write("!ENDTRNS\n")
        for entry in mye.entries:
            for i, line in enumerate(entry.lines):
                tag = "TRNS" if i == 0 else "SPL"
                memo = line.memo.replace("\t", " ")
                fh.write(
                    f"{tag}\tGENERAL JOURNAL\t{line.date}\t{acct(line.account_code)}\t"
                    f"{line.amount:.2f}\t{line.reference}\t{memo}\n"
                )
            fh.write("ENDTRNS\n")
    return path


# ----- editable round trip (unpack / pack) ----------------------------------


def unpack(mye: MyeFile, out_dir: str) -> List[str]:
    """Write the editable representation: accounts.csv, journal.csv,
    company.csv, Extract.inf and any extra archive members (e.g.
    BASLINK.TXT). Edit them in Excel or a text editor, then rebuild a
    .MYE with pack_dir()."""
    paths = export_csv(mye, out_dir)
    inf_path = os.path.join(out_dir, EXTRACT_INF)
    with open(inf_path, "wb") as fh:
        fh.write(mye.extract_inf)
    paths.append(inf_path)
    # Carry through extra members (BASLINK.TXT etc.) verbatim so a
    # round trip never loses BAS/GST or other side-car data.
    for name, payload in mye.extra_members.items():
        member_path = os.path.join(out_dir, name)
        with open(member_path, "wb") as fh:
            fh.write(payload)
        paths.append(member_path)
    return paths


def _read_csv(path: str) -> List[dict]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def pack_dir(in_dir: str) -> MyeFile:
    """Rebuild a MyeFile from a directory created by unpack() (after any
    edits). journal.csv rows are grouped into entries by the 'entry'
    column; the 'amount' column wins if present, otherwise debit/credit."""
    mye = MyeFile()

    company_path = os.path.join(in_dir, COMPANY_CSV)
    fields = {row["field"]: row["value"] for row in _read_csv(company_path)}
    company = ["", "", "", "", "", ""]
    for i, name in enumerate(COMPANY_FIELD_NAMES):
        company[i] = fields.get(name, "")
    # Any further positional fields (company_field_6, ...) if present.
    i = len(COMPANY_FIELD_NAMES)
    while f"company_field_{i}" in fields:
        company.append(fields[f"company_field_{i}"])
        i += 1
    mye.company_fields = company

    term = fields.get("meta.journal_line_terminator", "CRCRLF")
    mye.journal_line_terminator = "\r\n" if term == "CRLF" else "\r\r\n"
    mye.trailing_blank = fields.get("meta.trailing_blank", "yes").lower() != "no"

    for row in _read_csv(os.path.join(in_dir, ACCOUNTS_CSV)):
        mye.accounts.append(
            Account(code=row["account_code"].strip(), name=row["account_name"])
        )

    current_key = None
    entry = None
    for lineno, row in enumerate(_read_csv(os.path.join(in_dir, JOURNAL_CSV)), 2):
        amount_text = (row.get("amount") or "").strip()
        try:
            if amount_text:
                amount = Decimal(amount_text)
                amount_dp = len(amount_text.split(".", 1)[1]) if "." in amount_text else 0
            else:
                debit = Decimal((row.get("debit") or "0").strip() or "0")
                credit = Decimal((row.get("credit") or "0").strip() or "0")
                amount = debit - credit
                amount_dp = 4
        except InvalidOperation as exc:
            raise ValueError(f"journal.csv row {lineno}: bad amount") from exc
        key = (row.get("entry") or "").strip() or row.get("reference", "")
        if entry is None or key != current_key:
            entry = JournalEntry()
            mye.entries.append(entry)
            current_key = key
        entry.lines.append(
            JournalLine(
                date=row["date"].strip(),
                reference=row["reference"].strip(),
                account_code=row["account_code"].strip(),
                amount=amount,
                memo=row.get("memo", ""),
                amount_dp=amount_dp,
            )
        )

    inf_path = os.path.join(in_dir, EXTRACT_INF)
    if os.path.exists(inf_path):
        with open(inf_path, "rb") as fh:
            mye.extract_inf = fh.read()
    # Restore extra members: every file in the dir that isn't one of the
    # CSVs or Extract.inf (e.g. BASLINK.TXT).
    known = {COMPANY_CSV, ACCOUNTS_CSV, JOURNAL_CSV, TRIAL_BALANCE_CSV, EXTRACT_INF}
    for name in sorted(os.listdir(in_dir)):
        if name in known:
            continue
        full = os.path.join(in_dir, name)
        if os.path.isfile(full):
            with open(full, "rb") as fh:
                mye.extra_members[name] = fh.read()
    return mye
